"""Prepare a Cholec80 frame-extraction export as a portal dataset root.

This module does not register extraction ``run.json`` as an inference RunManifest
and does not fabricate predictions or metrics.
"""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal, cast

from openpyxl import load_workbook

from cvs_review_portal.domain.errors import InvalidIdentifierError
from cvs_review_portal.domain.identifiers import (
    validate_identifier,
    validate_relative_storage_key,
)
from cvs_review_portal.domain.models import (
    Catalog,
    CatalogDatasetRef,
    DatasetImage,
    DatasetVersion,
    ImageProvenance,
)
from cvs_review_portal.domain.types import parse_binary_label

_VIDEO_FILENAME_RE = re.compile(r"^video0*(\d+)\.mp4$", re.IGNORECASE)
_REQUIRED_XLSX_COLUMNS = (
    "video",
    "critical_view",
    "initial_minute",
    "initial_second",
    "final_minute",
    "final_second",
)


class Cholec80ImportError(RuntimeError):
    """Raised when Cholec80 import validation or preparation fails."""


@dataclass(frozen=True, slots=True)
class GroundTruthRow:
    video_number: int
    start_seconds: int
    end_seconds: int
    critical_view: Literal[0, 1]
    excel_row: int  # 1-based spreadsheet row for diagnostics only


@dataclass(frozen=True, slots=True)
class ManifestImage:
    image_id: str
    video_id: str
    segment_id: str
    frame_index: int
    timestamp: float
    path: str
    video_filename: str
    start_seconds: int
    end_seconds: int
    excel_row: int | None


@dataclass(frozen=True, slots=True)
class SegmentKey:
    video_number: int
    start_seconds: int
    end_seconds: int


@dataclass(frozen=True, slots=True)
class ExtractionVideo:
    video_id: str
    source_filenames: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ExtractionSegment:
    segment_id: str
    video_id: str
    video_filename: str
    start_seconds: int
    end_seconds: int
    exported_count: int


@dataclass(frozen=True, slots=True)
class ExtractionRunSummary:
    export_status: str
    succeeded_videos: int
    skipped_jobs: tuple[str, ...]
    declared_segments: int
    declared_images: int
    counts_segments_completed: int
    counts_images_exported: int
    succeeded_video_ids: frozenset[str]
    succeeded_filenames: frozenset[str]
    skipped_video_ids: frozenset[str]
    skipped_filenames: frozenset[str]
    videos_by_id: dict[str, ExtractionVideo]
    segments_by_id: dict[str, ExtractionSegment]


def _empty_str_list() -> list[str]:
    return []


def _empty_int_dict() -> dict[int, int]:
    return {}


@dataclass(slots=True)
class ImportReport:
    export_status: str
    succeeded_videos: int
    skipped_jobs: list[str] = field(default_factory=_empty_str_list)
    segment_count: int = 0
    image_count: int = 0
    ground_truth_segments: dict[int, int] = field(default_factory=_empty_int_dict)
    ground_truth_images: dict[int, int] = field(default_factory=_empty_int_dict)
    unused_xlsx_rows: int = 0
    repeated_frame_label_conflicts: list[str] = field(default_factory=_empty_str_list)
    dry_run: bool = False
    destination: Path | None = None


@dataclass(frozen=True, slots=True)
class PreparedImage:
    source_path: Path
    raw_image_key: str
    record: DatasetImage


def _as_nonnegative_int(value: object, *, field_name: str, excel_row: int) -> int:
    if type(value) is bool or value is None:
        raise Cholec80ImportError(
            f"Invalid {field_name} at XLSX row {excel_row}: missing or non-integer."
        )
    if type(value) is int:
        number = value
    elif type(value) is float:
        if not value.is_integer():
            raise Cholec80ImportError(
                f"Invalid {field_name} at XLSX row {excel_row}: non-integer value {value!r}."
            )
        number = int(value)
    else:
        raise Cholec80ImportError(
            f"Invalid {field_name} at XLSX row {excel_row}: "
            f"unsupported type {type(value).__name__}."
        )
    if number < 0:
        raise Cholec80ImportError(
            f"Invalid {field_name} at XLSX row {excel_row}: negative value {number}."
        )
    return number


def _as_binary_label(value: object, *, excel_row: int) -> Literal[0, 1]:
    if value is None:
        raise Cholec80ImportError(
            f"Missing critical_view at XLSX row {excel_row}; refusing to default to 0."
        )
    if type(value) is bool:
        raise Cholec80ImportError(
            f"Invalid critical_view at XLSX row {excel_row}: boolean is not allowed."
        )
    if type(value) is float:
        if value not in (0.0, 1.0) or not value.is_integer():
            raise Cholec80ImportError(f"Invalid critical_view at XLSX row {excel_row}: {value!r}.")
        value = int(value)
    try:
        return parse_binary_label(value)
    except ValueError as exc:
        raise Cholec80ImportError(
            f"Invalid critical_view at XLSX row {excel_row}: {value!r}."
        ) from exc


def video_number_from_filename(video_filename: str) -> int:
    match = _VIDEO_FILENAME_RE.fullmatch(video_filename.strip())
    if match is None:
        raise Cholec80ImportError(f"Unsupported video filename for join: {video_filename!r}.")
    return int(match.group(1))


def video_number_from_xlsx(value: object, *, excel_row: int) -> int:
    number = _as_nonnegative_int(value, field_name="video", excel_row=excel_row)
    if number < 1:
        raise Cholec80ImportError(f"Invalid video number at XLSX row {excel_row}: {number}.")
    return number


def build_display_id(
    *,
    dataset_id: str,
    dataset_version: str,
    video_filename: str,
    start_seconds: int,
    end_seconds: int,
    frame_index: int,
) -> str:
    stem = Path(video_filename).stem
    return f"{dataset_id}-{dataset_version}-{stem}-s{start_seconds}-e{end_seconds}-f{frame_index}"


def _try_segment_key_from_xlsx_cells(
    *,
    video: object,
    initial_minute: object,
    initial_second: object,
    final_minute: object,
    final_second: object,
    excel_row: int,
) -> SegmentKey | None:
    """Best-effort key for workbook row selection. Invalid/unrelated rows return None."""
    try:
        video_number = video_number_from_xlsx(video, excel_row=excel_row)
        start_seconds = _as_nonnegative_int(
            initial_minute, field_name="initial_minute", excel_row=excel_row
        ) * 60 + _as_nonnegative_int(
            initial_second, field_name="initial_second", excel_row=excel_row
        )
        end_seconds = _as_nonnegative_int(
            final_minute, field_name="final_minute", excel_row=excel_row
        ) * 60 + _as_nonnegative_int(final_second, field_name="final_second", excel_row=excel_row)
    except Cholec80ImportError:
        return None
    return SegmentKey(video_number, start_seconds, end_seconds)


def _strict_ground_truth_row(
    *,
    video: object,
    critical_view: object,
    initial_minute: object,
    initial_second: object,
    final_minute: object,
    final_second: object,
    excel_row: int,
) -> GroundTruthRow:
    video_number = video_number_from_xlsx(video, excel_row=excel_row)
    start_seconds = _as_nonnegative_int(
        initial_minute, field_name="initial_minute", excel_row=excel_row
    ) * 60 + _as_nonnegative_int(initial_second, field_name="initial_second", excel_row=excel_row)
    end_seconds = _as_nonnegative_int(
        final_minute, field_name="final_minute", excel_row=excel_row
    ) * 60 + _as_nonnegative_int(final_second, field_name="final_second", excel_row=excel_row)
    if end_seconds < start_seconds:
        raise Cholec80ImportError(
            f"Invalid interval at XLSX row {excel_row}: "
            f"end {end_seconds} < start {start_seconds}."
        )
    label = _as_binary_label(critical_view, excel_row=excel_row)
    return GroundTruthRow(
        video_number=video_number,
        start_seconds=start_seconds,
        end_seconds=end_seconds,
        critical_view=label,
        excel_row=excel_row,
    )


def index_ground_truth(rows: list[GroundTruthRow]) -> dict[SegmentKey, GroundTruthRow]:
    indexed: dict[SegmentKey, GroundTruthRow] = {}
    for row in rows:
        key = SegmentKey(row.video_number, row.start_seconds, row.end_seconds)
        if key in indexed:
            other = indexed[key]
            raise Cholec80ImportError(
                "Ambiguous ground-truth match: multiple XLSX rows share "
                f"video={row.video_number} start={row.start_seconds} end={row.end_seconds} "
                f"(rows {other.excel_row} and {row.excel_row})."
            )
        indexed[key] = row
    return indexed


def load_ground_truth_for_required_keys(
    xlsx_path: Path,
    required_keys: set[SegmentKey],
) -> tuple[dict[SegmentKey, GroundTruthRow], int]:
    """Validate workbook rows only for keys required by the export.

    Unrelated workbook rows (including invalid intervals / duplicate keys for
    excluded videos) are ignored and counted as unused. The owner's XLSX is
    never modified.
    """
    if not xlsx_path.is_file():
        raise Cholec80ImportError(f"Ground-truth workbook not found: {xlsx_path}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise Cholec80ImportError("Ground-truth workbook has no active sheet.")
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise Cholec80ImportError("Ground-truth workbook is empty.") from exc
        header = [str(cell).strip() if cell is not None else "" for cell in header_row]
        index = {name: position for position, name in enumerate(header)}
        missing = [name for name in _REQUIRED_XLSX_COLUMNS if name not in index]
        if missing:
            raise Cholec80ImportError(
                f"Ground-truth workbook missing required columns: {', '.join(missing)}."
            )

        selected: list[GroundTruthRow] = []
        unused_xlsx_rows = 0
        data_rows = 0
        for offset, row in enumerate(rows_iter, start=2):
            if all(cell is None or cell == "" for cell in row):
                continue
            data_rows += 1
            tentative = _try_segment_key_from_xlsx_cells(
                video=row[index["video"]],
                initial_minute=row[index["initial_minute"]],
                initial_second=row[index["initial_second"]],
                final_minute=row[index["final_minute"]],
                final_second=row[index["final_second"]],
                excel_row=offset,
            )
            if tentative is None or tentative not in required_keys:
                unused_xlsx_rows += 1
                continue
            selected.append(
                _strict_ground_truth_row(
                    video=row[index["video"]],
                    critical_view=row[index["critical_view"]],
                    initial_minute=row[index["initial_minute"]],
                    initial_second=row[index["initial_second"]],
                    final_minute=row[index["final_minute"]],
                    final_second=row[index["final_second"]],
                    excel_row=offset,
                )
            )

        if data_rows == 0:
            raise Cholec80ImportError("Ground-truth workbook contains no data rows.")

        gt_index = index_ground_truth(selected)
        missing_keys = required_keys - set(gt_index.keys())
        if missing_keys:
            sample = sorted(missing_keys, key=lambda key: (key.video_number, key.start_seconds))[0]
            raise Cholec80ImportError(
                "Missing ground-truth match for exported segment "
                f"video={sample.video_number} start={sample.start_seconds} "
                f"end={sample.end_seconds}."
            )
        return gt_index, unused_xlsx_rows
    finally:
        workbook.close()


def load_manifest_images(manifest_path: Path) -> list[ManifestImage]:
    if not manifest_path.is_file():
        raise Cholec80ImportError(f"manifest.jsonl not found: {manifest_path}")
    images: list[ManifestImage] = []
    with manifest_path.open(encoding="utf-8") as handle:
        for line_number, raw_line in enumerate(handle, start=1):
            line = raw_line.strip()
            if not line:
                continue
            try:
                raw_payload: object = json.loads(line)
            except json.JSONDecodeError as exc:
                raise Cholec80ImportError(
                    f"Invalid JSON on manifest.jsonl line {line_number}."
                ) from exc
            if not isinstance(raw_payload, dict):
                raise Cholec80ImportError(
                    f"manifest.jsonl line {line_number} must be a JSON object."
                )
            payload = cast(dict[str, Any], raw_payload)
            images.append(_parse_manifest_image(payload, line_number=line_number))
    if not images:
        raise Cholec80ImportError("manifest.jsonl contains no image records.")
    return images


def _require_str(payload: dict[str, Any], key: str, *, line_number: int) -> str:
    value = payload.get(key)
    if type(value) is not str or not value:
        raise Cholec80ImportError(
            f"manifest.jsonl line {line_number}: missing/invalid string field {key!r}."
        )
    return value


def _require_int(payload: dict[str, Any], key: str, *, line_number: int) -> int:
    value = payload.get(key)
    if type(value) is bool or value is None:
        raise Cholec80ImportError(
            f"manifest.jsonl line {line_number}: missing/invalid int field {key!r}."
        )
    if type(value) is int:
        return value
    if type(value) is float and value.is_integer():
        return int(value)
    raise Cholec80ImportError(
        f"manifest.jsonl line {line_number}: missing/invalid int field {key!r}."
    )


def _require_float(payload: dict[str, Any], key: str, *, line_number: int) -> float:
    value = payload.get(key)
    if type(value) is bool or value is None:
        raise Cholec80ImportError(
            f"manifest.jsonl line {line_number}: missing/invalid number field {key!r}."
        )
    if type(value) is int:
        return float(value)
    if type(value) is float:
        return value
    raise Cholec80ImportError(
        f"manifest.jsonl line {line_number}: missing/invalid number field {key!r}."
    )


def _parse_manifest_image(payload: dict[str, Any], *, line_number: int) -> ManifestImage:
    excel_row_raw = payload.get("excel_row")
    excel_row: int | None
    if excel_row_raw is None:
        excel_row = None
    else:
        excel_row = _require_int(payload, "excel_row", line_number=line_number)
    return ManifestImage(
        image_id=_require_str(payload, "image_id", line_number=line_number),
        video_id=_require_str(payload, "video_id", line_number=line_number),
        segment_id=_require_str(payload, "segment_id", line_number=line_number),
        frame_index=_require_int(payload, "frame_index", line_number=line_number),
        timestamp=_require_float(payload, "timestamp", line_number=line_number),
        path=_require_str(payload, "path", line_number=line_number),
        video_filename=_require_str(payload, "video_filename", line_number=line_number),
        start_seconds=_require_int(payload, "start_seconds", line_number=line_number),
        end_seconds=_require_int(payload, "end_seconds", line_number=line_number),
        excel_row=excel_row,
    )


def load_extraction_run(run_path: Path) -> dict[str, Any]:
    if not run_path.is_file():
        raise Cholec80ImportError(f"Extraction run.json not found: {run_path}")
    try:
        raw_payload: object = json.loads(run_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise Cholec80ImportError(f"Could not read extraction run.json: {run_path}") from exc
    if not isinstance(raw_payload, dict):
        raise Cholec80ImportError("Extraction run.json must be a JSON object.")
    return cast(dict[str, Any], raw_payload)


def _require_run_str(payload: dict[str, Any], key: str, *, context: str) -> str:
    value = payload.get(key)
    if type(value) is not str or not value:
        raise Cholec80ImportError(f"Extraction run.json {context}: missing/invalid {key!r}.")
    return value


def _require_run_int(payload: dict[str, Any], key: str, *, context: str) -> int:
    value = payload.get(key)
    if type(value) is bool or value is None:
        raise Cholec80ImportError(f"Extraction run.json {context}: missing/invalid {key!r}.")
    if type(value) is int:
        return value
    if type(value) is float and value.is_integer():
        return int(value)
    raise Cholec80ImportError(f"Extraction run.json {context}: missing/invalid {key!r}.")


def parse_extraction_run(run_payload: dict[str, Any]) -> ExtractionRunSummary:
    export_status = str(run_payload.get("status", "unknown"))
    jobs_raw = run_payload.get("jobs")
    if not isinstance(jobs_raw, list):
        raise Cholec80ImportError("Extraction run.json is missing a jobs list.")
    jobs = cast(list[object], jobs_raw)

    skipped_jobs: list[str] = []
    succeeded = 0
    declared_segments = 0
    declared_images = 0
    succeeded_video_ids: set[str] = set()
    succeeded_filenames: set[str] = set()
    skipped_video_ids: set[str] = set()
    skipped_filenames: set[str] = set()
    job_video_to_filename: dict[str, str] = {}

    for job_raw in jobs:
        if not isinstance(job_raw, dict):
            raise Cholec80ImportError("Extraction run.json contains a non-object job.")
        job = cast(dict[str, Any], job_raw)
        source = str(job.get("source_filename", "<unknown>"))
        status = str(job.get("status", "unknown"))
        video_id_raw = job.get("video_id")
        video_id = video_id_raw if type(video_id_raw) is str and video_id_raw else None
        if status == "succeeded":
            succeeded += 1
            declared_segments += int(job.get("segments_completed") or 0)
            declared_images += int(job.get("images_exported") or 0)
            if video_id is None:
                raise Cholec80ImportError(
                    f"Succeeded extraction job {source!r} is missing video_id."
                )
            succeeded_video_ids.add(video_id)
            succeeded_filenames.add(source)
            job_video_to_filename[video_id] = source
        elif status == "skipped":
            reason = ""
            issues_raw = job.get("issues")
            if isinstance(issues_raw, list) and issues_raw:
                issues = cast(list[object], issues_raw)
                first_raw = issues[0]
                if isinstance(first_raw, dict):
                    first = cast(dict[str, Any], first_raw)
                    reason = str(first.get("code") or first.get("message") or "")
            skipped_jobs.append(f"{source}" + (f" ({reason})" if reason else ""))
            skipped_filenames.add(source)
            if video_id is not None:
                skipped_video_ids.add(video_id)
        else:
            raise Cholec80ImportError(
                f"Extraction job {source!r} has unsupported status {status!r}."
            )

    if export_status not in {"completed", "partial"}:
        raise Cholec80ImportError(
            f"Unsupported extraction status {export_status!r}; refusing import."
        )
    if export_status == "partial" and succeeded == 0:
        raise Cholec80ImportError("Extraction status is partial with no succeeded jobs.")

    counts_raw = run_payload.get("counts")
    if not isinstance(counts_raw, dict):
        raise Cholec80ImportError("Extraction run.json is missing counts.")
    counts = cast(dict[str, Any], counts_raw)
    counts_segments = _require_run_int(counts, "segments_completed", context="counts")
    counts_images = _require_run_int(counts, "images_exported", context="counts")

    videos_raw = run_payload.get("videos")
    if not isinstance(videos_raw, list):
        raise Cholec80ImportError("Extraction run.json is missing videos[].")
    videos_by_id: dict[str, ExtractionVideo] = {}
    for index, video_raw in enumerate(cast(list[object], videos_raw)):
        if not isinstance(video_raw, dict):
            raise Cholec80ImportError(f"Extraction run.json videos[{index}] must be an object.")
        video_obj = cast(dict[str, Any], video_raw)
        video_id = _require_run_str(video_obj, "video_id", context=f"videos[{index}]")
        names_raw = video_obj.get("source_filenames")
        if not isinstance(names_raw, list) or not names_raw:
            raise Cholec80ImportError(
                f"Extraction run.json videos[{index}] is missing source_filenames."
            )
        names = tuple(str(name) for name in cast(list[object], names_raw))
        if video_id in videos_by_id:
            raise Cholec80ImportError(f"Duplicate video_id in run.json videos[]: {video_id}")
        videos_by_id[video_id] = ExtractionVideo(video_id=video_id, source_filenames=names)

    segments_raw = run_payload.get("segments")
    if not isinstance(segments_raw, list):
        raise Cholec80ImportError("Extraction run.json is missing segments[].")
    segments_by_id: dict[str, ExtractionSegment] = {}
    for index, segment_raw in enumerate(cast(list[object], segments_raw)):
        if not isinstance(segment_raw, dict):
            raise Cholec80ImportError(f"Extraction run.json segments[{index}] must be an object.")
        segment_obj = cast(dict[str, Any], segment_raw)
        segment_id = _require_run_str(segment_obj, "segment_id", context=f"segments[{index}]")
        if segment_id in segments_by_id:
            raise Cholec80ImportError(f"Duplicate segment_id in run.json segments[]: {segment_id}")
        segments_by_id[segment_id] = ExtractionSegment(
            segment_id=segment_id,
            video_id=_require_run_str(segment_obj, "video_id", context=f"segments[{index}]"),
            video_filename=_require_run_str(
                segment_obj, "video_filename", context=f"segments[{index}]"
            ),
            start_seconds=_require_run_int(
                segment_obj, "start_seconds", context=f"segments[{index}]"
            ),
            end_seconds=_require_run_int(segment_obj, "end_seconds", context=f"segments[{index}]"),
            exported_count=_require_run_int(
                segment_obj, "exported_count", context=f"segments[{index}]"
            ),
        )

    if counts_segments != declared_segments:
        raise Cholec80ImportError(
            "Extraction counts.segments_completed does not match succeeded-job totals: "
            f"counts={counts_segments} jobs_sum={declared_segments}."
        )
    if counts_images != declared_images:
        raise Cholec80ImportError(
            "Extraction counts.images_exported does not match succeeded-job totals: "
            f"counts={counts_images} jobs_sum={declared_images}."
        )
    if len(videos_by_id) != succeeded:
        raise Cholec80ImportError(
            "Extraction videos[] length does not match succeeded jobs: "
            f"videos={len(videos_by_id)} succeeded_jobs={succeeded}."
        )
    if len(segments_by_id) != counts_segments:
        raise Cholec80ImportError(
            "Extraction segments[] length does not match counts.segments_completed: "
            f"segments={len(segments_by_id)} counts={counts_segments}."
        )
    segment_exported_sum = sum(segment.exported_count for segment in segments_by_id.values())
    if segment_exported_sum != counts_images:
        raise Cholec80ImportError(
            "Sum of segments[].exported_count does not match counts.images_exported: "
            f"segments_sum={segment_exported_sum} counts={counts_images}."
        )

    for video_id, video in videos_by_id.items():
        if video_id in skipped_video_ids:
            raise Cholec80ImportError(
                f"Committed videos[] entry {video_id} belongs to a skipped job."
            )
        if video_id not in succeeded_video_ids:
            raise Cholec80ImportError(f"Committed videos[] entry {video_id} has no succeeded job.")
        expected_filename = job_video_to_filename.get(video_id)
        if expected_filename is None or expected_filename not in video.source_filenames:
            raise Cholec80ImportError(
                f"Video {video_id} source_filenames do not agree with succeeded job filename."
            )

    for segment in segments_by_id.values():
        if segment.video_id not in videos_by_id:
            raise Cholec80ImportError(
                f"Segment {segment.segment_id} references unknown video_id {segment.video_id}."
            )
        video = videos_by_id[segment.video_id]
        if segment.video_filename not in video.source_filenames:
            raise Cholec80ImportError(
                f"Segment {segment.segment_id} video_filename "
                f"{segment.video_filename!r} is absent from videos[]."
            )
        if segment.video_filename in skipped_filenames:
            raise Cholec80ImportError(
                f"Segment {segment.segment_id} uses skipped filename {segment.video_filename!r}."
            )

    return ExtractionRunSummary(
        export_status=export_status,
        succeeded_videos=succeeded,
        skipped_jobs=tuple(skipped_jobs),
        declared_segments=declared_segments,
        declared_images=declared_images,
        counts_segments_completed=counts_segments,
        counts_images_exported=counts_images,
        succeeded_video_ids=frozenset(succeeded_video_ids),
        succeeded_filenames=frozenset(succeeded_filenames),
        skipped_video_ids=frozenset(skipped_video_ids),
        skipped_filenames=frozenset(skipped_filenames),
        videos_by_id=videos_by_id,
        segments_by_id=segments_by_id,
    )


def reconcile_manifest_with_run(
    *,
    manifest_images: list[ManifestImage],
    run: ExtractionRunSummary,
) -> dict[str, SegmentKey]:
    """Verify committed membership between manifest.jsonl and run.json catalogs."""
    if len(manifest_images) != run.counts_images_exported:
        raise Cholec80ImportError(
            "Manifest image count does not match extraction run totals: "
            f"manifest={len(manifest_images)} counts.images_exported={run.counts_images_exported}."
        )

    seen_image_ids: set[str] = set()
    seen_paths: set[str] = set()
    segment_keys: dict[str, SegmentKey] = {}
    images_per_segment: Counter[str] = Counter()
    manifest_video_ids: set[str] = set()

    for image in manifest_images:
        if image.image_id in seen_image_ids:
            raise Cholec80ImportError(f"Duplicate manifest image_id: {image.image_id}")
        seen_image_ids.add(image.image_id)
        try:
            raw_image_key = validate_relative_storage_key(image.path, field="raw_image_key")
            validate_identifier(image.image_id, field="image_id")
            validate_identifier(image.video_id, field="video_id")
            validate_identifier(image.segment_id, field="segment_id")
        except InvalidIdentifierError as exc:
            raise Cholec80ImportError(str(exc)) from exc
        if raw_image_key in seen_paths:
            raise Cholec80ImportError(f"Duplicate manifest image path: {raw_image_key}")
        seen_paths.add(raw_image_key)

        if image.video_filename in run.skipped_filenames or image.video_id in run.skipped_video_ids:
            raise Cholec80ImportError(
                f"Manifest image {image.image_id} belongs to a skipped extraction job "
                f"({image.video_filename})."
            )
        if image.video_id not in run.succeeded_video_ids or image.video_id not in run.videos_by_id:
            raise Cholec80ImportError(
                f"Manifest image {image.image_id} video_id {image.video_id} "
                "is not a succeeded committed video."
            )
        video = run.videos_by_id[image.video_id]
        if image.video_filename not in video.source_filenames:
            raise Cholec80ImportError(
                f"Manifest image {image.image_id} video_filename "
                f"{image.video_filename!r} does not agree with videos[]."
            )
        if image.video_filename not in run.succeeded_filenames:
            raise Cholec80ImportError(
                f"Manifest image {image.image_id} uses non-succeeded filename "
                f"{image.video_filename!r}."
            )

        committed = run.segments_by_id.get(image.segment_id)
        if committed is None:
            raise Cholec80ImportError(
                f"Manifest segment {image.segment_id} is absent from run.json segments[]."
            )
        if (
            committed.video_id != image.video_id
            or committed.video_filename != image.video_filename
            or committed.start_seconds != image.start_seconds
            or committed.end_seconds != image.end_seconds
        ):
            raise Cholec80ImportError(
                f"Manifest segment {image.segment_id} does not agree with run.json "
                "video_id/filename/interval bounds."
            )

        video_number = video_number_from_filename(image.video_filename)
        key = SegmentKey(video_number, image.start_seconds, image.end_seconds)
        if image.segment_id in segment_keys and segment_keys[image.segment_id] != key:
            raise Cholec80ImportError(
                f"Segment {image.segment_id} maps to inconsistent interval bounds."
            )
        segment_keys[image.segment_id] = key
        images_per_segment[image.segment_id] += 1
        manifest_video_ids.add(image.video_id)

    if set(segment_keys) != set(run.segments_by_id):
        raise Cholec80ImportError(
            "Manifest segment membership does not match run.json segments[] exactly."
        )
    if manifest_video_ids != set(run.videos_by_id):
        raise Cholec80ImportError(
            "Manifest video membership does not match run.json videos[] exactly."
        )
    if len(segment_keys) != run.counts_segments_completed:
        raise Cholec80ImportError(
            "Unique manifest segment count does not match extraction run totals: "
            f"manifest_segments={len(segment_keys)} "
            f"counts.segments_completed={run.counts_segments_completed}."
        )

    for segment_id, committed in run.segments_by_id.items():
        actual = images_per_segment.get(segment_id, 0)
        if actual != committed.exported_count:
            raise Cholec80ImportError(
                f"Segment {segment_id} image count mismatch: "
                f"manifest={actual} exported_count={committed.exported_count}."
            )

    return segment_keys


def _assert_empty_destination(destination: Path) -> None:
    if not destination.exists():
        return
    if any(destination.rglob("*")):
        raise Cholec80ImportError(
            f"Refusing to write into nonempty destination: {destination}. "
            "Choose a new empty data root; existing owner data is left untouched."
        )


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def prepare_cholec80_dataset(
    *,
    export_root: Path,
    ground_truth_xlsx: Path,
    destination: Path,
    dataset_id: str,
    dataset_version: str,
    dry_run: bool = False,
    display_name: str | None = None,
    description: str | None = None,
) -> ImportReport:
    """Validate and optionally materialize a portal data root from a Cholec80 export."""
    export_root = export_root.resolve()
    ground_truth_xlsx = ground_truth_xlsx.resolve()
    destination = destination.resolve()

    try:
        validate_identifier(dataset_id, field="dataset_id")
        validate_identifier(dataset_version, field="dataset_version")
    except InvalidIdentifierError as exc:
        raise Cholec80ImportError(str(exc)) from exc

    if not export_root.is_dir():
        raise Cholec80ImportError(f"Export root not found: {export_root}")

    run_payload = load_extraction_run(export_root / "run.json")
    run = parse_extraction_run(run_payload)
    manifest_images = load_manifest_images(export_root / "manifest.jsonl")
    segment_keys = reconcile_manifest_with_run(manifest_images=manifest_images, run=run)

    required_keys = set(segment_keys.values())
    gt_index, unused_xlsx_rows = load_ground_truth_for_required_keys(
        ground_truth_xlsx, required_keys
    )

    prepared: list[PreparedImage] = []
    segment_labels: dict[SegmentKey, Literal[0, 1]] = {}
    frames_by_source: dict[tuple[str, int], set[Literal[0, 1]]] = defaultdict(set)

    for image in manifest_images:
        raw_image_key = validate_relative_storage_key(image.path, field="raw_image_key")
        key = segment_keys[image.segment_id]
        gt = gt_index[key]
        segment_labels[key] = gt.critical_view
        frames_by_source[(image.video_filename, image.frame_index)].add(gt.critical_view)

        source_path = (export_root / Path(*_posix_parts(raw_image_key))).resolve()
        try:
            source_path.relative_to(export_root)
        except ValueError as exc:
            raise Cholec80ImportError(f"Image path escapes export root: {raw_image_key}") from exc
        if not source_path.is_file():
            raise Cholec80ImportError(f"Missing source image file: {raw_image_key}")

        display_id = build_display_id(
            dataset_id=dataset_id,
            dataset_version=dataset_version,
            video_filename=image.video_filename,
            start_seconds=image.start_seconds,
            end_seconds=image.end_seconds,
            frame_index=image.frame_index,
        )
        provenance = ImageProvenance(
            video_id=image.video_id,
            video_filename=image.video_filename,
            segment_id=image.segment_id,
            start_seconds=image.start_seconds,
            end_seconds=image.end_seconds,
            frame_index=image.frame_index,
            timestamp=image.timestamp,
            display_id=display_id,
        )
        record = DatasetImage(
            image_id=image.image_id,
            raw_image_key=raw_image_key,
            ground_truth=gt.critical_view,
            comments=None,
            provenance=provenance,
        )
        prepared.append(
            PreparedImage(source_path=source_path, raw_image_key=raw_image_key, record=record)
        )

    conflict_messages: list[str] = []
    for (video_filename, frame_index), labels in sorted(frames_by_source.items()):
        if len(labels) > 1:
            conflict_messages.append(
                f"{video_filename} frame_index={frame_index} labels={sorted(labels)}"
            )

    dataset = DatasetVersion(
        schema_version="1",
        dataset_id=dataset_id,
        dataset_version=dataset_version,
        display_name=display_name or f"{dataset_id} {dataset_version}",
        description=description
        or (
            "Cholec80 CVS frame-export dataset prepared for portal review. "
            "No inference runs are published with this import."
        ),
        images=[item.record for item in prepared],
    )
    catalog = Catalog(
        schema_version="1",
        datasets=[
            CatalogDatasetRef(dataset_id=dataset_id, dataset_version=dataset_version),
        ],
        runs=[],
    )

    gt_segment_counts = Counter(segment_labels.values())
    gt_image_counts = Counter(item.record.ground_truth for item in prepared)
    report = ImportReport(
        export_status=run.export_status,
        succeeded_videos=run.succeeded_videos,
        skipped_jobs=list(run.skipped_jobs),
        segment_count=len(segment_keys),
        image_count=len(prepared),
        ground_truth_segments={0: gt_segment_counts.get(0, 0), 1: gt_segment_counts.get(1, 0)},
        ground_truth_images={0: gt_image_counts.get(0, 0), 1: gt_image_counts.get(1, 0)},
        unused_xlsx_rows=unused_xlsx_rows,
        repeated_frame_label_conflicts=conflict_messages,
        dry_run=dry_run,
        destination=None if dry_run else destination,
    )

    _assert_empty_destination(destination)
    if dry_run:
        return report

    destination.mkdir(parents=True, exist_ok=True)

    raw_root = destination / "raw-images" / dataset_id / dataset_version
    for item in prepared:
        target = raw_root.joinpath(*_posix_parts(item.raw_image_key))
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(item.source_path, target)

    dataset_path = (
        destination / "runs-output" / "datasets" / dataset_id / dataset_version / "dataset.json"
    )
    _write_json(dataset_path, dataset.model_dump(mode="json"))
    catalog_path = destination / "runs-output" / "catalog.json"
    _write_json(catalog_path, catalog.model_dump(mode="json"))
    return report


def _posix_parts(key: str) -> tuple[str, ...]:
    """Split a validated relative storage key into path parts."""
    return tuple(part for part in key.replace("\\", "/").split("/") if part)


def format_report(report: ImportReport) -> str:
    lines = [
        f"dry_run={report.dry_run}",
        f"export_status={report.export_status}",
        f"succeeded_videos={report.succeeded_videos}",
        f"skipped_jobs={len(report.skipped_jobs)}",
        *[f"  skipped: {item}" for item in report.skipped_jobs],
        f"segments={report.segment_count}",
        f"images={report.image_count}",
        f"gt_segments_0={report.ground_truth_segments.get(0, 0)}",
        f"gt_segments_1={report.ground_truth_segments.get(1, 0)}",
        f"gt_images_0={report.ground_truth_images.get(0, 0)}",
        f"gt_images_1={report.ground_truth_images.get(1, 0)}",
        f"unused_xlsx_rows={report.unused_xlsx_rows}",
        f"repeated_frame_label_conflicts={len(report.repeated_frame_label_conflicts)}",
    ]
    for item in report.repeated_frame_label_conflicts[:20]:
        lines.append(f"  conflict: {item}")
    if len(report.repeated_frame_label_conflicts) > 20:
        lines.append(f"  ... {len(report.repeated_frame_label_conflicts) - 20} more conflict(s)")
    if report.destination is not None:
        lines.append(f"destination={report.destination}")
    return "\n".join(lines)
